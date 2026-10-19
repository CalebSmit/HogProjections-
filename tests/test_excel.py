"""The Excel workbook must compute the same numbers as the Python model.

The workbook is recalculated with LibreOffice (headless). The test is skipped if
LibreOffice isn't installed; CI installs it on Linux.
"""

import shutil
import subprocess
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from hogproj.excel import COST_NAMES, SENS_CORN, SENS_HOG, build_workbook
from hogproj.inputs import InterestMode, LaborCostMode, MarketingCut, WeanPigPriceSource
from hogproj.model import project, sensitivity
from tests.test_model import make

SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")


# LibreOffice keeps cached formula results on load unless told to always recalculate.
RECALC_XCU = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry"
  xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<item oor:path="/org.openoffice.Office.Calc/Formula/Load">
  <prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop>
</item>
</oor:items>
"""


def recalc(xlsx_bytes: bytes, tmp_path: Path) -> openpyxl.Workbook:
    src = tmp_path / "in.xlsx"
    src.write_bytes(xlsx_bytes)
    out = tmp_path / "out"
    profile = tmp_path / "lo_profile"
    (profile / "user").mkdir(parents=True)
    (profile / "user" / "registrymodifications.xcu").write_text(RECALC_XCU)
    assert SOFFICE
    subprocess.run(
        [
            SOFFICE,
            f"-env:UserInstallation={profile.as_uri()}",
            "--headless",
            "--calc",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(out),
            str(src),
        ],
        check=True,
        capture_output=True,
        timeout=120,
    )
    return openpyxl.load_workbook(out / "in.xlsx", data_only=True)


def named(wb: openpyxl.Workbook, name: str):
    """Value of a single-cell named range (results are read by name, not position)."""
    dn = wb.defined_names[name]
    sheet, coord = next(iter(dn.destinations))
    return wb[sheet][coord.replace("$", "")].value  # type: ignore[union-attr]


def named_grid(wb: openpyxl.Workbook, name: str) -> list[list[float]]:
    sheet, coord = next(iter(wb.defined_names[name].destinations))
    cells = wb[sheet][coord.replace("$", "")]
    return [[float(c.value) for c in row] for row in cells]  # type: ignore[arg-type, union-attr]


def costs_by_line(wb: openpyxl.Workbook) -> dict[str, float]:
    return {line: float(named(wb, n) or 0) for line, n in COST_NAMES.items()}


needs_lo = pytest.mark.skipif(SOFFICE is None, reason="LibreOffice not installed")


@needs_lo
@pytest.mark.parametrize(
    "overrides",
    [
        {},
        {
            "checkoff_pct": 0.35,
            "interest_rate_pct": 8.0,
            "dead_feed_share": 0.3,
            "wean_pig_source": WeanPigPriceSource.MANUAL,
            "wean_pig_manual_usd_hd": 55.0,
            "labor_mode": LaborCostMode.PER_PIG_PLACED,
            "labor_usd_pig_placed": 3.0,
            "lean_hog_futures_cwt": {"2026-06": 100.0, "2026-07": 104.0},
            "hog_basis_cwt_by_period": {10: -2.0, 12: 1.5},
            "marketing": [
                MarketingCut(days_after_placement=140, pct_of_marketed=40, live_weight_lb=280),
                MarketingCut(days_after_placement=165, pct_of_marketed=60, live_weight_lb=290),
            ],
        },
    ],
)
def test_workbook_matches_model(tmp_path, overrides):
    p = make(**overrides)
    check_workbook(p, [p.location], tmp_path)


@needs_lo
def test_workbook_edge_cases(tmp_path):
    """Stale copy of the selected location in the list, an unselected location sorted first,
    a 0% cut with no futures price, and a Dec 20 sale rolling to next year's Feb contract."""
    p = make(
        placement_date=date(2026, 7, 23),
        marketing=[
            MarketingCut(days_after_placement=150, pct_of_marketed=100, live_weight_lb=285),
            MarketingCut(days_after_placement=200, pct_of_marketed=0, live_weight_lb=300),
        ],
        lean_hog_futures_cwt={"2027-02": 90.0},
        hog_basis_cwt_by_period={24: 4.0},
    )
    assert project(p).marketing.loc[0, "contract"] == "2027-02"
    stale = p.location.model_copy(update={"corn_basis_usd_bu": -9.0})
    other = p.location.model_copy(update={"name": "AAA first", "ddgs_usd_ton": 999.0})
    check_workbook(p, [other, stale], tmp_path)


@needs_lo
def test_default_projection_workbook(tmp_path):
    from hogproj.data.defaults import default_projection, load_locations

    locs = load_locations()
    p = default_projection(locs[1], placement_date=date(2026, 10, 19))
    check_workbook(p, locs, tmp_path)


@needs_lo
def test_workbook_matches_crush_layout_inputs(tmp_path):
    """Every new input (adder, hedges, carry, $/pig mortality and interest, G&A, other, debt)."""
    from tests.test_crush_parity import app_projection

    p = app_projection()
    check_workbook(p, [p.location], tmp_path)


@needs_lo
def test_workbook_rate_interest_and_pct_mortality_with_new_lines(tmp_path):
    p = make(
        ga_usd_hd=2.0,
        other_usd_hd_sold=-1.5,
        wean_pig_adder_usd_hd=6.0,
        corn_carry_usd_bu=0.2,
        corn_hedged_pct=30,
        corn_hedged_usd_bu=4.1,
        hog_hedged_pct=50,
        hog_hedged_cwt=105.0,
        interest_rate_pct=7.0,
        principal_usd_hd=3.0,
        term_interest_usd_hd=1.0,
    )
    check_workbook(p, [p.location], tmp_path)


def check_workbook(p, locations, tmp_path):
    expected = project(p)
    wb = recalc(build_workbook(p, locations), tmp_path)

    got = costs_by_line(wb)
    for line, total in expected.costs.set_index("line")["total_usd"].items():
        assert got[str(line)] == pytest.approx(total, rel=1e-9, abs=1e-6), line

    s = expected.summary
    assert named(wb, "KpiNet") == pytest.approx(s["net_usd"])
    assert named(wb, "KpiNetHd") == pytest.approx(s["net_usd_per_head_sold"])
    assert named(wb, "KpiRevenue") == pytest.approx(s["gross_revenue_usd"])
    assert named(wb, "KpiDebt") == pytest.approx(s["debt_service_usd"])
    assert named(wb, "KpiNetAfterDebtHd") == pytest.approx(s["net_after_debt_usd_per_head_sold"])
    assert named(wb, "KpiBreakeven") == pytest.approx(s["breakeven_cwt_carcass"])
    assert named(wb, "KpiAvgPrice") == pytest.approx(s["avg_price_cwt_carcass"])
    assert named(wb, "KpiCostHd") == pytest.approx(s["cost_usd_per_head_sold"])
    assert named(wb, "KpiFeedConv") == pytest.approx(s["feed_conversion"])
    grid = named_grid(wb, "SensGrid")
    want = sensitivity(p, SENS_CORN, SENS_HOG).to_numpy()
    assert [v for row in grid for v in row] == pytest.approx(want.ravel().tolist())


@needs_lo
@pytest.mark.parametrize("corn_is_cash", [False, True])
def test_weekly_crush_sheet_matches_crush_table(tmp_path, corn_is_cash):
    from hogproj.crush import CrushSpec, Week, crush_table
    from tests.test_crush_parity import app_projection

    p = app_projection().model_copy(
        update={"interest_mode": InterestMode.RATE, "interest_rate_pct": 7.0}
    )
    weeks = [
        Week(date(2026, 9, 13), 40.0, 4.5, 330.0, 160.0, 80.0, -3.0),
        Week(date(2026, 9, 6), 45.0, 4.3, 320.0, 170.0, 78.0, 1.0),
        Week(date(2026, 12, 20), 30.0, 4.6, 340.0, 150.0, 85.0, 0.5),  # Dec 15+ -> next Feb
    ]
    spec = CrushSpec(weeks, 1, 160, 290.0, corn_is_cash)
    expected = crush_table(p, weeks, 1, 160, 290.0, corn_is_cash)
    ws = recalc(build_workbook(p, [p.location], crush=spec), tmp_path)["Weekly Crush"]
    labels = {str(ws.cell(r, 1).value): r for r in range(1, 120) if ws.cell(r, 1).value}
    for j, col in enumerate(expected.columns):
        c = j + 2
        assert ws.cell(labels["CME contract"], c).value == expected.loc["CME contract", col]
        for label in (
            "Revenue $/pig",
            "Feed $/pig",
            "Cost of production $/pig",
            "Margin $/pig",
            "Margin after debt service $/pig",
            "Corn $/bu paid",
        ):
            got = ws.cell(labels[label], c).value
            assert got == pytest.approx(float(expected.loc[label, col])), (label, col)


@needs_lo
def test_scenario_name_reaches_the_dashboard(tmp_path):
    p = make()
    wb = recalc(build_workbook(p, [p.location], scenario='Smith\'s "north" site'), tmp_path)
    assert str(wb["Dashboard"]["B2"].value).startswith('Smith\'s "north" site')
    (tmp_path / "b").mkdir()
    blank = recalc(build_workbook(p, [p.location]), tmp_path / "b")
    assert str(blank["Dashboard"]["B2"].value).startswith(p.location.name)
