import io
from datetime import date

import openpyxl
import pandas as pd
import pytest

from hogproj.export import build_report
from hogproj.outlook import Plan, project, summary

START = date(2026, 10, 19)
NO_SEASON = pd.DataFrame(columns=["month", "factor", "years"])


def test_report_opens_and_ties_to_projection():
    plan = Plan(seasonal=False)
    weekly = project(plan, NO_SEASON, START)
    data = build_report(plan, weekly, START, "Sample data.", "Losing money at today's price.")
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Summary", "Monthly", "Costs", "Weekly"]

    s = wb["Summary"]
    values = {s.cell(r, 1).value: s.cell(r, 2).value for r in range(1, s.max_row + 1)}
    assert values["Profit, next 12 months"] == pytest.approx(summary(weekly)["profit_usd"])
    assert values["Sows today"] == 1000
    assert s["A4"].value == "Losing money at today's price."

    w = wb["Weekly"]
    assert w.max_row == len(weekly) + 1
    assert w.cell(2, 8).value == pytest.approx(weekly["margin_usd"].iloc[0])
    assert wb["Monthly"].max_row == 25  # header + 24 months


def test_brief_workbook_has_every_section():
    from hogproj.export import build_brief

    data = build_brief(
        START,
        "Hogs down $3.11 on the week.",
        [("Lean Hog Index", 73.40, -3.11, -0.23, "$/cwt"), ("Iowa corn", 4.24, None, 0.09, "$/bu")],
        [("Placing pigs", "A pig placed today projects -$7.38.")],
        {"Pigs on feed": pd.DataFrame({"contract": ["Apr 2027"], "margin": [-13.5]}),
         "By year": pd.DataFrame({"year": [2025, 2026], "index": [94.4, float("nan")]})},
    )  # fmt: skip
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Brief", "Pigs on feed", "By year"]
    b = wb["Brief"]
    assert b["A4"].value == "Hogs down $3.11 on the week."
    cells = [c.value for row in b.iter_rows() for c in row]
    assert "Lean Hog Index" in cells and 73.40 in cells
    assert "A pig placed today projects -$7.38." in cells
    assert wb["By year"]["B3"].value is None  # NaN written blank
