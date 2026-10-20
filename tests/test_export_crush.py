"""The live crush workbook: sheets, named inputs, formulas, and cached values that match the app."""

import io
from datetime import date

import openpyxl
import pandas as pd
import pytest

from hogproj.crush import CrushInputs, margin_usd
from hogproj.export_crush import ContractRow, Placement, build_crush_workbook, margins

C = CrushInputs(
    pig_usd_per_head=20.08,
    adder_usd_per_head=7.0,
    corn_usd_per_bu=4.31,
    principal_usd=3.1,
    term_interest_usd=1.3,
)
K = [
    ContractRow("Apr 2027", date(2027, 4, 14), 70.50),
    ContractRow("Feb 2027", date(2027, 2, 12), 65.50, hedged=0.5, hedge_price=72.0),
]
P = [
    Placement(
        date(2026, 10, 20), date(2026, 10, 20), date(2027, 3, 29), "Apr 2027", 20.08, 4.31, 1000
    ),
    Placement(
        date(2026, 9, 18), date(2026, 9, 14), date(2027, 2, 21), "Feb 2027", 28.71, 4.40, 1000
    ),
]


def _book():
    hist = pd.DataFrame(
        {
            "placed": pd.to_datetime(["2025-01-06"]),
            "sold": pd.to_datetime(["2025-06-15"]),
            "pig_usd": [60.0],
            "worth_usd": [70.0],
            "hog_usd_per_cwt": [100.0],
            "corn_usd_per_bu": [4.1],
            "cost_usd": [150.0],
            "margin_usd": [10.0],
        }
    )
    usda = pd.DataFrame({"week_ending": pd.to_datetime(["2026-10-16"]), "weaned_pig": [20.08]})
    data = build_crush_workbook(C, P, K, (0.0, 4.0), hist, usda, date(2026, 10, 20))
    return (
        openpyxl.load_workbook(io.BytesIO(data)),
        openpyxl.load_workbook(io.BytesIO(data), data_only=True),
    )


def test_sheets_and_named_inputs():
    wb, _ = _book()
    assert wb.sheetnames == ["Read me", "Inputs", "Contracts", "Crush", "History", "USDA data"]
    names = set(wb.defined_names.keys())
    assert {"Adder", "DaysOnFeed", "Basis", "CornLb", "ContractTable", "CornHedgePct"} <= names


def test_margin_cells_are_formulas_with_app_values():
    wb, vals = _book()
    crush, cached = wb["Crush"], vals["Crush"]
    assert str(crush["B16"].value).startswith("=")  # margin / pig is live
    assert "VLOOKUP" in str(crush["B11"].value)
    # Column B: Apr 2027 unhedged -> the app's own margin (debt included in app cost).
    assert cached["B18"].value == pytest.approx(margin_usd(C, 70.50, pig=20.08, corn=4.31))
    # Column C: Feb 2027 half hedged at $72 -> effective $68.75.
    assert cached["C11"].value == pytest.approx(0.5 * 72 + 0.5 * 65.5)
    assert vals["Contracts"]["F3"].value == pytest.approx(68.75)


def test_totals_and_corn_hedge():
    _, vals = _book()
    m = [margins(C, p, {k.label: k for k in K}, (0.0, 4.0))["after_debt"] for p in P]
    assert vals["Crush"]["D20"].value == pytest.approx(1000 * sum(m))
    hedged = margins(C, P[0], {k.label: k for k in K}, (1.0, 3.50))
    plain = margins(C, P[0], {k.label: k for k in K}, (0.0, 3.50))
    assert hedged["after_debt"] - plain["after_debt"] == pytest.approx(
        C.corn_lb * (4.31 - 3.5) / 56
    )
